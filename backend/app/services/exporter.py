import csv
import io
import json
from typing import List

from ..models import JobPost


def export_jobs_to_csv_bytes(jobs: List[JobPost]) -> bytes:
    """Export jobs list to UTF-8 with BOM CSV bytes (opens cleanly in Microsoft Excel with Vietnamese accents)."""
    output = io.StringIO()
    writer = csv.writer(output)
    writer.writerow(
        [
            "ID",
            "Tiêu đề (Title)",
            "Công ty (Company)",
            "Mức lương gốc (Raw Salary)",
            "Lương quy đổi (Triệu VNĐ/tháng)",
            "Địa điểm (Location)",
            "Hình thức (Work Mode)",
            "Cấp bậc (Seniority)",
            "Kỹ năng (Tech Stack)",
            "HR Email",
            "HR SĐT",
            "Tên HR",
            "Quy mô công ty",
            "Nguồn (Domain)",
            "Trạng thái JD",
            "URL Ứng tuyển",
            "Ngày lấy",
        ]
    )

    for j in jobs:
        skills_str = ""
        try:
            skills_str = ", ".join(json.loads(j.skills_json or "[]"))
        except Exception:
            skills_str = j.skills_json or ""

        writer.writerow(
            [
                j.id,
                j.title,
                j.company_name,
                j.salary_raw,
                j.salary_vnd_normalized if not j.is_negotiable else "Thỏa thuận",
                j.location_city,
                j.work_mode,
                j.seniority,
                skills_str,
                j.hr_email or "",
                j.hr_phone or "",
                j.hr_name or "",
                j.company_size or "",
                j.source_domain,
                j.deep_crawl_status,
                j.url,
                j.created_at.strftime("%Y-%m-%d %H:%M") if j.created_at else "",
            ]
        )

    # Prepend UTF-8 BOM so Excel on Windows renders Vietnamese characters properly
    return b"\xef\xbb\xbf" + output.getvalue().encode("utf-8")


def export_jobs_to_xlsx_bytes(jobs: List[JobPost]) -> bytes:
    """Export jobs list to a styled Excel (.xlsx) workbook in memory."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "Crawled Jobs"

    headers = [
        "ID",
        "Tiêu đề (Title)",
        "Công ty (Company)",
        "Mức lương gốc",
        "Quy đổi (Triệu VNĐ)",
        "Địa điểm",
        "Hình thức",
        "Cấp bậc",
        "Kỹ năng (Skills)",
        "HR Email",
        "HR SĐT",
        "Tên HR",
        "Quy mô công ty",
        "Nguồn",
        "URL Ứng tuyển",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True)
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for j in jobs:
        skills_str = ""
        try:
            skills_str = ", ".join(json.loads(j.skills_json or "[]"))
        except Exception:
            skills_str = ""

        ws.append(
            [
                j.id,
                j.title,
                j.company_name,
                j.salary_raw,
                j.salary_vnd_normalized if not j.is_negotiable else 0.0,
                j.location_city,
                j.work_mode,
                j.seniority,
                skills_str,
                j.hr_email or "",
                j.hr_phone or "",
                j.hr_name or "",
                j.company_size or "",
                j.source_domain,
                j.url,
            ]
        )

    col_widths = [8, 38, 28, 22, 18, 18, 14, 15, 35, 25, 16, 18, 18, 18, 50]
    for i, width in enumerate(col_widths, start=1):
        col_letter = ws.cell(row=1, column=i).column_letter
        ws.column_dimensions[col_letter].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
